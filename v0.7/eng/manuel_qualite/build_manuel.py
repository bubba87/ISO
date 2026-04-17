#!/usr/bin/env python3
"""Build a consolidated Quality Manual (Word + PDF) from all MQ_*.md files."""

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Cm

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, KeepTogether,
)

SCRIPT_DIR = Path(__file__).resolve().parent
OUT_DOCX = SCRIPT_DIR / "Manuel_Qualite_v0.7_EN.docx"
OUT_PDF = SCRIPT_DIR / "Manuel_Qualite_v0.7_EN.pdf"

NAVY = colors.HexColor("#1a365d")
DARK = colors.HexColor("#2d3748")


# ---------- shared inline formatter ----------
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\n]+\*|`[^`\n]+`)")


def md_inline_to_rl(text: str) -> str:
    """Convert markdown inline markers to ReportLab mini-HTML."""
    def repl(m):
        t = m.group(0)
        if t.startswith("**"):
            return f"<b>{escape_rl(t[2:-2])}</b>"
        if t.startswith("*"):
            return f"<i>{escape_rl(t[1:-1])}</i>"
        if t.startswith("`"):
            return f'<font name="Courier">{escape_rl(t[1:-1])}</font>'
        return escape_rl(t)

    out = []
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            out.append(escape_rl(text[pos:m.start()]))
        out.append(repl(m))
        pos = m.end()
    if pos < len(text):
        out.append(escape_rl(text[pos:]))
    return "".join(out)


def escape_rl(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ---------- PDF (ReportLab) ----------
def build_pdf(md_files):
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"],
                          fontName="Helvetica", fontSize=10.5, leading=14,
                          textColor=colors.HexColor("#222222"), spaceAfter=6)
    h1 = ParagraphStyle("h1", parent=styles["Heading1"],
                        fontName="Helvetica-Bold", fontSize=18, leading=22,
                        textColor=NAVY, spaceBefore=10, spaceAfter=10,
                        borderPadding=4, borderWidth=0, borderColor=NAVY)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"],
                        fontName="Helvetica-Bold", fontSize=13, leading=17,
                        textColor=NAVY, spaceBefore=10, spaceAfter=6)
    h3 = ParagraphStyle("h3", parent=styles["Heading3"],
                        fontName="Helvetica-Bold", fontSize=11.5, leading=15,
                        textColor=DARK, spaceBefore=8, spaceAfter=4)
    h4 = ParagraphStyle("h4", parent=h3, fontSize=10.5)
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=18,
                            bulletIndent=6, spaceAfter=2)
    number = ParagraphStyle("number", parent=body, leftIndent=22,
                            bulletIndent=6, spaceAfter=2)
    cover_title = ParagraphStyle("cover_title", parent=body,
                                 fontName="Helvetica-Bold", fontSize=30,
                                 leading=36, alignment=TA_CENTER, textColor=NAVY,
                                 spaceAfter=20)
    cover_line = ParagraphStyle("cover_line", parent=body, fontSize=14,
                                leading=20, alignment=TA_CENTER,
                                textColor=colors.HexColor("#4a5568"))

    heading_styles = {1: h1, 2: h2, 3: h3, 4: h4, 5: h4, 6: h4}

    story = []

    # Cover
    story.append(Spacer(1, 6 * cm))
    story.append(Paragraph("Quality Manual", cover_title))
    story.append(Paragraph("<b>Plus Sarl</b>", cover_line))
    story.append(Spacer(1, 0.4 * cm))
    story.append(Paragraph("ISO 9001:2015", cover_line))
    story.append(Spacer(1, 0.4 * cm))
    story.append(Paragraph("Version 0.7 — English", cover_line))
    story.append(PageBreak())

    for idx, md_file in enumerate(md_files):
        if idx > 0:
            story.append(PageBreak())
        render_md_to_flowables(md_file.read_text(encoding="utf-8"), story,
                                body, bullet, number, heading_styles)

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawCentredString(A4[0] / 2, 1.2 * cm,
                                 f"Plus Sarl — Quality Manual v0.7 — Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(OUT_PDF), pagesize=A4,
                            leftMargin=1.8 * cm, rightMargin=1.8 * cm,
                            topMargin=2 * cm, bottomMargin=2.2 * cm,
                            title="Quality Manual v0.7 — Plus Sarl",
                            author="Plus Sarl")
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    size_kb = OUT_PDF.stat().st_size / 1024
    print(f"PDF  : {OUT_PDF.name} ({size_kb:.0f} KB)")


def render_md_to_flowables(md_text, story, body, bullet, number, heading_styles):
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue

        if re.match(r"^---+$", line):
            story.append(Spacer(1, 0.3 * cm))
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            story.append(Paragraph(md_inline_to_rl(m.group(2).strip()),
                                   heading_styles[level]))
            i += 1
            continue

        # Pipe table
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[i + 1]):
            tbl_lines = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                tbl_lines.append(lines[i])
                i += 1
            rows = []
            for tl in tbl_lines:
                if re.match(r"^\s*\|?\s*:?-+", tl):
                    continue
                cells = [c.strip() for c in tl.strip().strip("|").split("|")]
                rows.append(cells)
            if rows:
                ncols = max(len(r) for r in rows)
                data = []
                for r_idx, row in enumerate(rows):
                    padded = row + [""] * (ncols - len(row))
                    data.append([Paragraph(md_inline_to_rl(c), body) for c in padded])
                tbl = Table(data, repeatRows=1, hAlign="LEFT")
                tbl.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                    ("TOPPADDING", (0, 0), (-1, 0), 6),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e0")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ]))
                story.append(tbl)
                story.append(Spacer(1, 0.3 * cm))
            continue

        # Bullet list
        if re.match(r"^[-*]\s+", line):
            while i < len(lines) and re.match(r"^[-*]\s+", lines[i].rstrip()):
                item = re.sub(r"^[-*]\s+", "", lines[i].rstrip())
                story.append(Paragraph(md_inline_to_rl(item), bullet,
                                        bulletText="•"))
                i += 1
            story.append(Spacer(1, 0.15 * cm))
            continue

        # Numbered list
        if re.match(r"^\d+\.\s+", line):
            n = 1
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].rstrip()):
                item = re.sub(r"^\d+\.\s+", "", lines[i].rstrip())
                story.append(Paragraph(md_inline_to_rl(item), number,
                                        bulletText=f"{n}."))
                n += 1
                i += 1
            story.append(Spacer(1, 0.15 * cm))
            continue

        # Paragraph
        para_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt:
                break
            if re.match(r"^(#{1,6}\s|[-*]\s|\d+\.\s|\|)", nxt):
                break
            if re.match(r"^---+$", nxt):
                break
            para_lines.append(nxt)
            i += 1
        story.append(Paragraph(md_inline_to_rl(" ".join(para_lines)), body))


# ---------- DOCX (python-docx) ----------
def add_inline_docx(paragraph, text: str):
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_md_to_docx(doc, md_text):
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue

        if re.match(r"^---+$", line):
            doc.add_paragraph().add_run("_" * 60)
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).strip(), level=min(level, 4))
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[i + 1]):
            tbl_lines = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                tbl_lines.append(lines[i])
                i += 1
            rows = []
            for tl in tbl_lines:
                if re.match(r"^\s*\|?\s*:?-+", tl):
                    continue
                cells = [c.strip() for c in tl.strip().strip("|").split("|")]
                rows.append(cells)
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = "Light Grid Accent 1"
                for r_idx, row in enumerate(rows):
                    for c_idx in range(ncols):
                        cell_text = row[c_idx] if c_idx < len(row) else ""
                        cell = table.rows[r_idx].cells[c_idx]
                        cell.text = ""
                        add_inline_docx(cell.paragraphs[0], cell_text)
            continue

        if re.match(r"^[-*]\s+", line):
            while i < len(lines) and re.match(r"^[-*]\s+", lines[i].rstrip()):
                item = re.sub(r"^[-*]\s+", "", lines[i].rstrip())
                p = doc.add_paragraph(style="List Bullet")
                add_inline_docx(p, item)
                i += 1
            continue

        if re.match(r"^\d+\.\s+", line):
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].rstrip()):
                item = re.sub(r"^\d+\.\s+", "", lines[i].rstrip())
                p = doc.add_paragraph(style="List Number")
                add_inline_docx(p, item)
                i += 1
            continue

        para_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt:
                break
            if re.match(r"^(#{1,6}\s|[-*]\s|\d+\.\s|\|)", nxt):
                break
            if re.match(r"^---+$", nxt):
                break
            para_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        add_inline_docx(p, " ".join(para_lines))


def build_docx(md_files):
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    title = doc.add_paragraph()
    title.alignment = 1
    run = title.add_run("Quality Manual")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    for line in ("Plus Sarl", "ISO 9001:2015", "Version 0.7 — English"):
        p = doc.add_paragraph(line)
        p.alignment = 1
        for r in p.runs:
            r.font.size = Pt(14)
    doc.add_page_break()

    for idx, md_file in enumerate(md_files):
        if idx > 0:
            doc.add_page_break()
        add_md_to_docx(doc, md_file.read_text(encoding="utf-8"))

    doc.save(OUT_DOCX)
    size_kb = OUT_DOCX.stat().st_size / 1024
    print(f"DOCX : {OUT_DOCX.name} ({size_kb:.0f} KB)")


def main():
    md_files = sorted(SCRIPT_DIR.glob("MQ_*.md"))
    print(f"Combining {len(md_files)} chapters...")
    build_pdf(md_files)
    build_docx(md_files)


if __name__ == "__main__":
    main()
