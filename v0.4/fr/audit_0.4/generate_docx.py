#!/usr/bin/env python3
"""
Générateur DOCX du dossier d'audit ISO 9001:2015 — Plus Sarl
Combine MQ-001 + INDEX en un seul fichier .docx pour transmission à SQS.
Utilise python-docx pour un rendu professionnel.
"""

import re
import os
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

# Couleurs
BLUE_DARK = RGBColor(0x1a, 0x36, 0x5d)
BLUE_HEADER = RGBColor(0x1a, 0x36, 0x5d)
BLUE_LIGHT = "D6E4F0"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x66, 0x66, 0x66)
BLACK = RGBColor(0x22, 0x22, 0x22)


def set_cell_shading(cell, color_hex):
    """Set background color of a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def add_formatted_run(paragraph, text, bold=False, italic=False, size=None, color=None, font_name=None):
    """Add a formatted run to a paragraph."""
    run = paragraph.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if font_name:
        run.font.name = font_name
    return run


def setup_styles(doc):
    """Configure document styles."""
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10)
    font.color.rgb = BLACK
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.3

    for level, (size, spacing_before) in enumerate(
        [(20, 24), (15, 18), (12, 14), (11, 10)], start=1
    ):
        heading_style = doc.styles[f'Heading {level}']
        heading_style.font.name = 'Calibri'
        heading_style.font.size = Pt(size)
        heading_style.font.color.rgb = BLUE_DARK
        heading_style.font.bold = True
        heading_style.paragraph_format.space_before = Pt(spacing_before)
        heading_style.paragraph_format.space_after = Pt(6)


def add_cover_page(doc):
    """Add professional cover page."""
    # Add empty paragraphs for spacing
    for _ in range(4):
        doc.add_paragraph()

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formatted_run(p, "DOSSIER D'AUDIT", bold=True, size=28, color=BLUE_DARK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formatted_run(p, "ISO 9001:2015", bold=True, size=24, color=BLUE_DARK)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formatted_run(p, "Plus Sarl", bold=True, size=20, color=BLUE_DARK)

    doc.add_paragraph()

    # Info table
    info = [
        ("Société", "Plus Sarl"),
        ("Adresse", "Route de Montet 11, 1588 Cudrefin, Suisse"),
        ("N. IDE", "CH-645.4.101.228-7"),
        ("Gérante", "Roxane Wicky"),
        ("Organisme de certification", "SQS"),
        ("Version", "0.4 — Version de travail"),
        ("Date", "20/02/2026"),
    ]

    table = doc.add_table(rows=len(info), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, (key, val) in enumerate(info):
        cell_key = table.cell(i, 0)
        cell_val = table.cell(i, 1)
        p_key = cell_key.paragraphs[0]
        p_key.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_formatted_run(p_key, key, bold=True, size=10, color=BLUE_DARK)
        p_val = cell_val.paragraphs[0]
        add_formatted_run(p_val, val, size=10)
        set_cell_shading(cell_key, BLUE_LIGHT)

    # Set column widths
    for row in table.rows:
        row.cells[0].width = Cm(6)
        row.cells[1].width = Cm(8)

    # Footer on cover page
    for _ in range(6):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formatted_run(p, "Partie 1 — Manuel Qualité (MQ-001)", size=10, color=GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formatted_run(p, "Partie 2 — Index du Dossier d'Audit", size=10, color=GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formatted_run(p, "Document préparé pour l'organisme de certification SQS", size=9, color=GRAY, italic=True)

    doc.add_page_break()


def parse_markdown_line(text):
    """Parse inline markdown (bold, italic, links) and return list of (text, bold, italic) tuples."""
    parts = []
    i = 0
    while i < len(text):
        # Bold: **text**
        m = re.match(r'\*\*(.+?)\*\*', text[i:])
        if m:
            parts.append((m.group(1), True, False))
            i += m.end()
            continue
        # Italic: *text*
        m = re.match(r'\*(.+?)\*', text[i:])
        if m:
            parts.append((m.group(1), False, True))
            i += m.end()
            continue
        # Link: [text](url) → just text
        m = re.match(r'\[(.+?)\]\(.+?\)', text[i:])
        if m:
            parts.append((m.group(1), False, False))
            i += m.end()
            continue
        # Inline code: `text`
        m = re.match(r'`(.+?)`', text[i:])
        if m:
            parts.append((m.group(1), False, False))
            i += m.end()
            continue
        # Regular character
        parts.append((text[i], False, False))
        i += 1

    # Merge consecutive same-style parts
    merged = []
    for text_part, bold, italic in parts:
        if merged and merged[-1][1] == bold and merged[-1][2] == italic:
            merged[-1] = (merged[-1][0] + text_part, bold, italic)
        else:
            merged.append((text_part, bold, italic))

    return merged


def add_rich_paragraph(doc, text, style='Normal', bold_all=False):
    """Add a paragraph with inline markdown formatting."""
    # Remove emoji patterns
    text = re.sub(r'🟢|🔴|🟡|🔵|🆕|:red_circle:|:yellow_circle:|:green_circle:|:blue_circle:', '', text)
    text = text.strip()

    if not text:
        return doc.add_paragraph()

    p = doc.add_paragraph(style=style)
    parts = parse_markdown_line(text)
    for part_text, bold, italic in parts:
        run = p.add_run(part_text)
        run.bold = bold or bold_all
        run.italic = italic
        run.font.name = 'Calibri'
    return p


def add_table_from_rows(doc, header_row, data_rows):
    """Add a formatted table to the document."""
    if not header_row:
        return

    cols = len(header_row)
    rows_count = 1 + len(data_rows)
    table = doc.add_table(rows=rows_count, cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    for j, cell_text in enumerate(header_row):
        cell = table.cell(0, j)
        cell_text = re.sub(r'🟢|🔴|🟡|🔵|🆕|:red_circle:|:yellow_circle:|:green_circle:|:blue_circle:', '', cell_text).strip()
        p = cell.paragraphs[0]
        parts = parse_markdown_line(cell_text)
        for part_text, bold, italic in parts:
            run = p.add_run(part_text)
            run.bold = True
            run.italic = italic
            run.font.name = 'Calibri'
            run.font.size = Pt(9)
            run.font.color.rgb = WHITE
        set_cell_shading(cell, "1A365D")

    # Data rows
    for i, row_data in enumerate(data_rows):
        for j in range(min(cols, len(row_data))):
            cell = table.cell(i + 1, j)
            cell_text = row_data[j]
            cell_text = re.sub(r'🟢|🔴|🟡|🔵|🆕|:red_circle:|:yellow_circle:|:green_circle:|:blue_circle:', '', cell_text).strip()
            p = cell.paragraphs[0]
            parts = parse_markdown_line(cell_text)
            for part_text, bold, italic in parts:
                run = p.add_run(part_text)
                run.bold = bold
                run.italic = italic
                run.font.name = 'Calibri'
                run.font.size = Pt(9)
            # Alternate row shading
            if i % 2 == 1:
                set_cell_shading(cell, "F0F4F8")

    doc.add_paragraph()  # spacing after table


def process_markdown_file(doc, filepath, is_part2=False):
    """Process a markdown file and add its content to the document."""
    with open(filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    in_code_block = False
    in_table = False
    table_header = None
    table_rows = []
    in_blockquote = False
    blockquote_lines = []

    while i < len(lines):
        line = lines[i].rstrip('\n')

        # Code blocks
        if line.strip().startswith('```'):
            if in_code_block:
                in_code_block = False
                i += 1
                continue
            else:
                in_code_block = True
                i += 1
                continue

        if in_code_block:
            # Add as preformatted text
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.name = 'Courier New'
            run.font.size = Pt(8)
            run.font.color.rgb = GRAY
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            i += 1
            continue

        # Flush table if we're leaving table context
        if in_table and not line.strip().startswith('|'):
            add_table_from_rows(doc, table_header, table_rows)
            in_table = False
            table_header = None
            table_rows = []

        # Flush blockquote if leaving blockquote context
        if in_blockquote and not line.strip().startswith('>'):
            bq_text = ' '.join(blockquote_lines)
            p = add_rich_paragraph(doc, bq_text)
            p.paragraph_format.left_indent = Cm(1)
            p.runs[0].italic = True if p.runs else None
            in_blockquote = False
            blockquote_lines = []

        stripped = line.strip()

        # Empty line
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped == '---' or stripped == '***':
            # Skip, we use headings for structure
            i += 1
            continue

        # Headings
        heading_match = re.match(r'^(#{1,4})\s+(.+)', stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            # Clean up text
            text = re.sub(r'🟢|🔴|🟡|🔵|🆕', '', text).strip()
            doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        # Table rows
        if stripped.startswith('|') and stripped.endswith('|'):
            cells = [c.strip() for c in stripped.split('|')[1:-1]]

            # Check if it's a separator row (|---|---|)
            if all(re.match(r'^[-:]+$', c) for c in cells):
                i += 1
                continue

            if not in_table:
                in_table = True
                table_header = cells
            else:
                table_rows.append(cells)
            i += 1
            continue

        # Blockquote
        if stripped.startswith('>'):
            bq_text = re.sub(r'^>\s*', '', stripped)
            if not in_blockquote:
                in_blockquote = True
                blockquote_lines = [bq_text]
            else:
                blockquote_lines.append(bq_text)
            i += 1
            continue

        # Checkbox items
        if stripped.startswith('- ['):
            text = re.sub(r'^- \[.\]\s*', '', stripped)
            add_rich_paragraph(doc, f"  {text}")
            i += 1
            continue

        # List items (bullet)
        list_match = re.match(r'^(\s*)-\s+(.+)', line)
        if list_match:
            indent = len(list_match.group(1))
            text = list_match.group(2)
            p = add_rich_paragraph(doc, text)
            p.style = 'List Bullet'
            if indent >= 2:
                p.paragraph_format.left_indent = Cm(1.5)
            i += 1
            continue

        # Numbered list
        num_match = re.match(r'^(\s*)\d+\.\s+(.+)', line)
        if num_match:
            text = num_match.group(2)
            p = add_rich_paragraph(doc, text)
            p.style = 'List Number'
            i += 1
            continue

        # Regular paragraph
        add_rich_paragraph(doc, stripped)
        i += 1

    # Flush remaining table
    if in_table:
        add_table_from_rows(doc, table_header, table_rows)

    # Flush remaining blockquote
    if in_blockquote:
        bq_text = ' '.join(blockquote_lines)
        p = add_rich_paragraph(doc, bq_text)
        p.paragraph_format.left_indent = Cm(1)


def main():
    doc = Document()

    # Page setup
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    setup_styles(doc)

    # === Cover page ===
    add_cover_page(doc)

    # === Part 1: Manuel Qualité ===
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formatted_run(p, "PARTIE 1", bold=True, size=14, color=BLUE_DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formatted_run(p, "Manuel Qualité — MQ-001", bold=True, size=18, color=BLUE_DARK)
    doc.add_paragraph()

    mq_path = os.path.join(SCRIPT_DIR, 'MQ-001_Manuel_Qualite.md')
    process_markdown_file(doc, mq_path)

    doc.add_page_break()

    # === Part 2: Index Dossier Audit ===
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formatted_run(p, "PARTIE 2", bold=True, size=14, color=BLUE_DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formatted_run(p, "Index du Dossier d'Audit", bold=True, size=18, color=BLUE_DARK)
    doc.add_paragraph()

    index_path = os.path.join(SCRIPT_DIR, 'INDEX_DOSSIER_AUDIT.md')
    process_markdown_file(doc, index_path, is_part2=True)

    # === Save ===
    output_path = os.path.join(SCRIPT_DIR, 'DOSSIER_AUDIT_ISO9001_Plus_Sarl_v0.4.docx')
    doc.save(output_path)

    size_kb = os.path.getsize(output_path) / 1024
    print(f"Document généré : {output_path}")
    print(f"Taille : {size_kb:.0f} Ko")


if __name__ == '__main__':
    main()
